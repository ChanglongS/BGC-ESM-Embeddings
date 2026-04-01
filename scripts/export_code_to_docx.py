import docx
from docx.shared import Pt
from docx.oxml.ns import qn
import os

def add_code_to_doc(doc, filepath):
    filename = os.path.basename(filepath)
    doc.add_heading(filename, level=2)
    
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            lines = f.readlines()
    except Exception as e:
        doc.add_paragraph(f"Error reading file: {e}")
        return

    # Create a single paragraph for the code block to keep it together
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    # Process line by line to ensure formatting
    for line in lines:
        run = p.add_run(line)
        # Use a monospaced font to ensure indentation aligns perfectly
        run.font.name = 'Courier New'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
        run.font.size = Pt(9)

def main():
    doc = docx.Document()
    doc.add_heading('BGC-DETR 核心脚本源代码', 0)
    
    files = [
        "/share/org/BGI/bgi_suncl/project/BGC-DETR/scripts/run_end_to_end_prediction.py",
        "/share/org/BGI/bgi_suncl/project/BGC-DETR/run_prodigal_for_genomes.py",
        "/share/org/BGI/bgi_suncl/project/BGC-DETR/process_prediction_data_fixed.py",
        "/share/org/BGI/bgi_suncl/project/BGC-DETR/validate_final.py",
        "/share/org/BGI/bgi_suncl/project/BGC-DETR/scripts/export_filtered_predictions.py"
    ]
    
    for f in files:
        if os.path.exists(f):
            print(f"Processing {f}...")
            add_code_to_doc(doc, f)
        else:
            print(f"Warning: File not found {f}")
            
    output_path = "/share/org/BGI/bgi_suncl/project/BGC-DETR/docs/BGC_DETR_核心代码.docx"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully saved to {output_path}")

if __name__ == "__main__":
    main()
